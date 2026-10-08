import re
from pathlib import Path
from typing import Protocol

from docx import Document
from pypdf import PdfReader

from paperhub.core.models import ParsedDocument
from paperhub.errors import PaperHubError


class Parser(Protocol):
    file_types: list[str]

    def parse(self, path: Path) -> ParsedDocument: ...


def metadata(text: str, fallback: str = "") -> ParsedDocument:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    title_match = re.search(r"\\title\s*\{([^}]+)\}", text)
    title = (
        title_match[1]
        if title_match
        else next(
            (line.lstrip("# ") for line in lines if not line.startswith(("---", "\\", "!"))),
            fallback,
        )
    )
    abstract = re.search(
        r"(?:^|\n)\s*(?:#{1,6}\s*)?(?:Abstract|摘要)\s*[:：]?\s*\n?"
        r"(.*?)(?=\n\s*(?:#{1,6}\s+|(?:\d+[. ]+)?Introduction\b|引言|Keywords\b|关键词)|\Z)",
        text,
        re.I | re.S,
    )
    tex_abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", text, re.S)
    doi = re.search(r"\b10\.\d{4,9}/[-._;()/:A-Za-z0-9]+", text, re.I)
    year = re.search(r"\b(?:19|20)\d{2}\b", text[:6000])
    authors = re.search(r"(?:^|\n)(?:Authors?|作者)\s*[:：]\s*([^\n]+)", text, re.I)
    tex_authors = re.search(r"\\author\s*\{([^}]+)\}", text)
    author_text = authors[1] if authors else tex_authors[1] if tex_authors else ""
    keywords = re.search(r"(?:^|\n)(?:Keywords?|关键词)\s*[:：]\s*([^\n]+)", text, re.I)
    return ParsedDocument(
        text=text,
        title=title[:500],
        authors=[a.strip() for a in re.split(r";|,|，|、|\\and", author_text) if a.strip()],
        year=int(year[0]) if year else None,
        abstract=(tex_abstract[1] if tex_abstract else abstract[1] if abstract else "").strip()[
            :20000
        ],
        keywords=[w.strip() for w in re.split(r"[,;，；]", keywords[1]) if w.strip()]
        if keywords
        else [],
        doi=doi[0].rstrip(".,;)") if doi else "",
        language="zh" if re.search(r"[\u4e00-\u9fff]", text[:2000]) else "en",
    )


class TextParser:
    file_types = ["md", "tex", "html"]

    def parse(self, path: Path) -> ParsedDocument:
        return metadata(path.read_text(encoding="utf-8-sig"), path.stem)


class DocxParser:
    file_types = ["docx"]

    def parse(self, path: Path) -> ParsedDocument:
        doc = Document(str(path))
        text = "\n\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(
            " | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows
        )
        result = metadata(text, path.stem)
        if doc.core_properties.title:
            result.title = doc.core_properties.title
        if not result.authors and doc.core_properties.author:
            result.authors = [doc.core_properties.author]
        return result


class PdfParser:
    file_types = ["pdf"]

    def parse(self, path: Path) -> ParsedDocument:
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise PaperHubError("E_PARSE_FAILED", "加密 PDF 无法读取 / Encrypted PDF")
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        if not text.strip():
            raise PaperHubError(
                "E_OCR_REQUIRED",
                "PDF 无可提取文本，需要 OCR / OCR required",
                "使用 OCRmyPDF 后重新扫描",
            )
        result = metadata(text, path.stem)
        if reader.metadata:
            pdf_title = reader.metadata.title
            if pdf_title and not pdf_title.lower().startswith(("untitled", "microsoft")):
                result.title = pdf_title
            if reader.metadata.author and not result.authors:
                result.authors = [reader.metadata.author]
        return result


class ParserRegistry:
    def __init__(self) -> None:
        self.parsers: dict[str, Parser] = {}
        for parser in (TextParser(), DocxParser(), PdfParser()):
            self.register(parser)

    def register(self, parser: Parser) -> None:
        for extension in parser.file_types:
            self.parsers[extension] = parser

    def parse(self, path: Path) -> ParsedDocument:
        parser = self.parsers.get(path.suffix.lower().lstrip("."))
        if not parser:
            raise PaperHubError("E_FORMAT_UNSUPPORTED", "文件格式不支持 / Unsupported format")
        try:
            return parser.parse(path)
        except PaperHubError:
            raise
        except Exception as exc:
            raise PaperHubError(
                "E_PARSE_FAILED", f"解析失败 / Parse failed ({type(exc).__name__})"
            ) from None
