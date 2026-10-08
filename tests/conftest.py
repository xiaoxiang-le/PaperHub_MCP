import asyncio
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from paperhub.app import Application
from paperhub.config import Config
from paperhub.core.models import CostEstimate, TranslateRequest, TranslateResponse
from paperhub.errors import PaperHubError


def make_pdf(path: Path, text: str) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    content = DecodedStreamObject()
    lines = text.splitlines()
    commands = ["BT /F1 12 Tf 50 740 Td"]
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        commands.append(f"({escaped}) Tj 0 -16 Td")
    content.set_data(("\n".join(commands) + "\nET").encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.add_metadata({"/Title": lines[0], "/Author": "Ada Lovelace"})
    with path.open("wb") as stream:
        writer.write(stream)


@pytest.fixture
def app(tmp_path: Path) -> Application:
    papers = tmp_path / "papers"
    papers.mkdir()
    config = Config.model_validate(
        {
            "data_dir": tmp_path / "state",
            "library": {"paths": [papers]},
            "convert": {"output_dir": tmp_path / "output"},
            "translate": {"chunk_chars": 500, "max_retries": 0, "timeout": 2},
        }
    )
    return Application(config)


@pytest.fixture
def library(app: Application) -> Application:
    root = app.guard.roots[0]
    (root / "attention.md").write_text(
        "# Efficient Transformer Attention\n\nAuthors: Ada; Bob\nYear: 2024\n\n"
        "## Abstract\nFast attention reduces transformer inference latency.\n\n"
        "Keywords: attention, transformer\nDOI: 10.1234/attention\n\n## Introduction\n"
        "Equation $E=mc^2$ and [1, 2]. Figure 1.\n```python\nx = 1\n```",
        encoding="utf-8",
    )
    (root / "graph.tex").write_text(
        r"\title{Graph Neural Networks}"
        + "\n"
        + r"\author{Carol\and Dan}"
        + "\n2023\n"
        + r"\begin{abstract}Graph message passing for molecular prediction.\end{abstract}",
        encoding="utf-8",
    )
    doc = Document()
    doc.core_properties.title = "Protein Structure Prediction"
    doc.add_paragraph("Protein Structure Prediction")
    doc.add_paragraph("Abstract\nProtein folding models predict molecular structure.")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Protein"
    doc.save(root / "protein.docx")
    make_pdf(root / "robot.pdf", "Robot Motion Planning\nAbstract\nRobot control and navigation.")
    app.library.scan(str(root), True, None)
    return app


class MockBackend:
    name = "mock"

    def __init__(self) -> None:
        self.calls: list[str] = []
        self.fail_on: str | None = None
        self.corrupt = False
        self.delay = 0.0

    def estimate_cost(self, tokens: int) -> CostEstimate:
        return CostEstimate(
            input_tokens=tokens, estimated_output_tokens=tokens, estimated_usd=(0, 0)
        )

    async def translate(self, req: TranslateRequest) -> TranslateResponse:
        self.calls.append(req.text)
        await asyncio.sleep(self.delay)
        if self.fail_on and self.fail_on in req.text:
            raise PaperHubError("E_BACKEND_FAILED", "mock failed")
        return TranslateResponse(
            text=req.text.replace("PH_", "BROKEN_") if self.corrupt else req.text
        )


def confirmation(call: Any) -> str:
    with pytest.raises(PaperHubError) as error:
        call()
    assert error.value.code == "E_CONFIRM_REQUIRED"
    return error.value.details["confirm_token"]


async def wait_task(app: Application, task_id: str) -> dict[str, Any]:
    for _ in range(1000):
        state = app.tasks.status(task_id)
        if state["status"] not in ("pending", "running"):
            return state
        await asyncio.sleep(0.01)
    raise AssertionError("task timed out")
